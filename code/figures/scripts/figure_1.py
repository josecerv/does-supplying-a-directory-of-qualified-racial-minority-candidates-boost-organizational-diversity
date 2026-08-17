"""
Figure 1 (Katy-feedback rebuild, round 5) — horizontal participant flow.

Round-5 (Jose, 2026-06-25): the 2->3 randomization is drawn as Katy's SHUFFLE
mark from her reference figure -- two strands that CROSS (the randomize
metaphor) and whose ends are the arrows into Treatment / Control. Drawn as
vector curves (custom geometry), not an icon-in-a-circle plus arrows, so there
is no "arrows on arrows" doubling. (Round 4 had a plain CONSORT fork.)


Round-4 = round-3 neutral palette + the verified fixes from the multi-agent
accuracy/design audit (8 agents, 2026-06-25), checked against the manuscript's
own Figure 1 caption (gdoc 1Ffgl..., docx 20260622-...-klm):

  MUST-FIX (accuracy):
   * Condition chips now show the actual manipulation: Treatment = a CURATED
     directory of qualified underrepresented racial minority (URM) faculty at
     peer departments; Control = links to peer-department websites that
     organizers SEARCH ON THEIR OWN. (URM expanded once, at first use.)
   * Step 4 names the real dependent variable: the SHARE of presenting speakers
     who were URM, and separately Black and Hispanic, in the 1,686 continuing
     seminar series (drops the misleading "all"; ties to Figure 2).
   * Step 3 names the randomization UNIT: departments (a cluster RCT).
   * Step 1 frames the encouragement to diversify as common to BOTH arms.
  SHOULD-FIX (polish):
   * Condition chips deepened + Outcome card softened to #333 so the two
     conditions are the salient color (color reserved for conditions).
   * The load-bearing randomization fork is darker/heavier than the decorative
     next-step chevrons. Gap rhythm rebalanced. Larger captions.
  DECLINED (audit guidance, keep it clean): per-arm n, 1:1 ratio, stratification
   line, 568-department / 146-university counts, attrition funnel -- the
   journal-typeset caption carries exact quantities.

Numbers verified: 1,586 recipients (1,019 organizers + 567 chairs) responsible
for 1,881 STEM seminar series at 146 R1 universities; departments randomized
(stratified by size); 1,686 series (89.6%) continued in 2024-25.

No em-dashes / semicolons. Instrument is a DIRECTORY (not database).
Builds figure_1.pptx (python-pptx); figure_1.pdf/.png are exported from PowerPoint
(see render_figure1.ps1). Canonical source: 08_manuscript/figure1/build_figure1_katy.py.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pathlib import Path
import os

OUT = Path(__file__).resolve().parent
RAND_STYLE = os.environ.get("RAND_STYLE", "shuffle")   # shuffle | dice | rnode
EMU_IN = 914400

# ── palette (GPT-5.5 consult + audit refinements) ────────────────────
# Structure == TRUE-neutral grays (zero saturation). Color is reserved
# ONLY for the two conditions, which are the salient elements.
NEUT_DARK = "242424"     # title / header text, step-4 number
CARD_DARK = "333333"     # Outcome endpoint card fill (softened from #242424)
NEUT_MED  = "737373"     # step number circles + step-1/2 icon badges
NEUT_LINE = "D9D9D9"     # card borders
CHEV = "8A8A8A"          # decorative next-step chevrons
FORK = "5F5F5F"          # load-bearing randomization fork (darker, heavier)
TX_DARK = "414141"       # caption / body text
TX_MUTE = "737373"
PILL_FILL = "EFEFEF"
CARD = "FFFFFF"
# Conditions == the ONLY color, matching Figure 2's left column.
TRT = "832040"; TRT_FILL = "ECD6DC"; TRT_BORDER = "D8B2BE"   # wine red
CTL = "2E5984"; CTL_FILL = "DAE4EF"; CTL_BORDER = "B5C6D6"   # steel blue
ICON = "Segoe MDL2 Assets"; SANS = "Segoe UI"
G_MAIL = chr(0xE166); G_KEY = chr(0xE192)
G_PEOPLE = chr(0xE716); G_DIR = chr(0xE779); G_LIST = chr(0xE71D)


def _solid(s, h): s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(h)
def _line(s, h, w=0.75):
    if h is None: s.line.fill.background()
    else: s.line.color.rgb = RGBColor.from_string(h); s.line.width = Pt(w)
def _no_shadow(s):
    try: s.shadow.inherit = False
    except Exception: pass
def soft_shadow(shape, blur=0.05, dist=0.03, color="93A6C2", alpha_pct=42):
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')): spPr.remove(el)
    eff = spPr.makeelement(qn('a:effectLst'), {})
    sh = spPr.makeelement(qn('a:outerShdw'), {'blurRad': str(int(blur*EMU_IN)),
        'dist': str(int(dist*EMU_IN)), 'dir': '5400000', 'rotWithShape': '0'})
    clr = spPr.makeelement(qn('a:srgbClr'), {'val': color})
    a = spPr.makeelement(qn('a:alpha'), {'val': str(int(alpha_pct*1000))})
    clr.append(a); sh.append(clr); eff.append(sh); spPr.append(eff)
def rrect(slide, x, y, w, h, fill, line=NEUT_LINE, line_w=0.75, radius=0.10, shadow=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try: sp.adjustments[0] = radius
    except Exception: pass
    _solid(sp, fill); _line(sp, line, line_w); _no_shadow(sp)
    if shadow: soft_shadow(sp)
    sp.text_frame.paragraphs[0].text = ""; return sp
def oval(slide, x, y, w, h, fill, line=None, line_w=0.75):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    _solid(sp, fill); _line(sp, line, line_w); _no_shadow(sp); return sp
def chevron(slide, cx, cy, w=0.26, h=0.38, fill=CHEV):
    sp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(cx-w/2), Inches(cy-h/2), Inches(w), Inches(h))
    _solid(sp, fill); _line(sp, None); _no_shadow(sp); return sp
def seg(slide, x1, y1, x2, y2, color, w=1.5, arrow=False):
    """Straight connector with round caps/joins; optional arrowhead at (x2,y2)."""
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = RGBColor.from_string(color); cn.line.width = Pt(w)
    ln = cn.line._get_or_add_ln(); ln.set('cap', 'rnd')
    ln.append(ln.makeelement(qn('a:round'), {}))
    if arrow:
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    _no_shadow(cn); return cn
def _bez(p0, c1, c2, p3, n=32):
    """Sample a cubic bezier into n+1 points (inches)."""
    out = []
    for i in range(n + 1):
        t = i / n; m = 1 - t
        out.append((m*m*m*p0[0] + 3*m*m*t*c1[0] + 3*m*t*t*c2[0] + t*t*t*p3[0],
                    m*m*m*p0[1] + 3*m*m*t*c1[1] + 3*m*t*t*c2[1] + t*t*t*p3[1]))
    return out
def curve_arrow(slide, pts, color, w=2.0):
    """Open custom-geometry polyline (a sampled curve) with an arrowhead at its end."""
    xs = [p[0] for p in pts]; ys = [p[1] for p in pts]
    L = min(xs); T = min(ys); Wd = max(1e-3, max(xs)-L); Ht = max(1e-3, max(ys)-T)
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(T), Inches(Wd), Inches(Ht))
    sp.fill.background(); _no_shadow(sp)
    sp.line.color.rgb = RGBColor.from_string(color); sp.line.width = Pt(w)
    ln = sp.line._get_or_add_ln(); ln.set('cap', 'rnd')
    ln.append(ln.makeelement(qn('a:round'), {}))
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'lg'}))
    spPr = sp._element.spPr
    for pg in spPr.findall(qn('a:prstGeom')): spPr.remove(pg)
    We = int(Wd*EMU_IN); He = int(Ht*EMU_IN)
    cg = spPr.makeelement(qn('a:custGeom'), {})
    for tag in ('a:avLst', 'a:gdLst', 'a:ahLst'):
        cg.append(cg.makeelement(qn(tag), {}))
    cg.append(cg.makeelement(qn('a:rect'), {'l': '0', 't': '0', 'r': str(We), 'b': str(He)}))
    pl = cg.makeelement(qn('a:pathLst'), {})
    path = cg.makeelement(qn('a:path'), {'w': str(We), 'h': str(He), 'fill': 'none'})
    def _pt(p):
        e = path.makeelement(qn('a:pt'), {'x': str(int((p[0]-L)*EMU_IN)), 'y': str(int((p[1]-T)*EMU_IN))})
        return e
    mv = path.makeelement(qn('a:moveTo'), {}); mv.append(_pt(pts[0])); path.append(mv)
    for p in pts[1:]:
        lt = path.makeelement(qn('a:lnTo'), {}); lt.append(_pt(p)); path.append(lt)
    pl.append(path); cg.append(pl)
    spPr.find(qn('a:xfrm')).addnext(cg)
    return sp
def text(slide, x, y, w, h, lines, *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = Emu(0); tf.margin_right = Emu(0); tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    first = True
    for para in lines:
        runs, pdict = (para if isinstance(para, tuple) else (para, {}))
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = pdict.get('align', align)
        if 'space_after' in pdict: p.space_after = Pt(pdict['space_after'])
        if 'space_before' in pdict: p.space_before = Pt(pdict['space_before'])
        if 'line_spacing' in pdict: p.line_spacing = pdict['line_spacing']
        for s, d in runs:
            r = p.add_run(); r.text = s; f = r.font
            f.size = Pt(d.get('size', 11)); f.bold = d.get('bold', False); f.italic = d.get('italic', False)
            f.name = d.get('font', SANS); f.color.rgb = RGBColor.from_string(d.get('color', TX_DARK))
    return tb
def badge(slide, cx, cy_top, glyph, *, d=0.82, isize=27, fill=NEUT_MED, gcolor="FFFFFF", shadow=True):
    o = oval(slide, cx-d/2, cy_top, d, d, fill, line=None)
    if shadow: soft_shadow(o, blur=0.045, dist=0.025, color="9AA2AC", alpha_pct=34)
    text(slide, cx-d/2, cy_top, d, d, [[(glyph, {'font': ICON, 'size': isize, 'color': gcolor})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
def num_circle(slide, cx, cy, n, *, d=0.42, fill=NEUT_MED, tcol="FFFFFF", ring=None):
    o = oval(slide, cx-d/2, cy-d/2, d, d, fill, line=ring, line_w=1.4)
    soft_shadow(o, blur=0.035, dist=0.02, color="9AA2AC", alpha_pct=30)
    text(slide, cx-d/2, cy-d/2, d, d, [[(str(n), {'size': 15, 'bold': True, 'color': tcol})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
def pill(slide, cx, y, label, *, fill=PILL_FILL, tx=NEUT_DARK):
    pw = 0.24 + 0.090*len(label); ph = 0.32
    rrect(slide, cx-pw/2, y, pw, ph, fill, line=None, radius=0.5)
    text(slide, cx-pw/2, y, pw, ph, [[(label, {'size': 11, 'bold': True, 'color': tx})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ── geometry ─────────────────────────────────────────────────────────
PAD = 0.18; DN = 0.42
TITLE_SZ = 14; CAP_SZ = 13


def header(slide, x, y0, w, n, title_lines, dark=False, size=TITLE_SZ,
           num_fill=NEUT_MED, num_tcol="FFFFFF"):
    num_circle(slide, x+PAD+DN/2, y0+0.14+DN/2, n, fill=num_fill, tcol=num_tcol,
               ring=("FFFFFF" if dark else None))
    tcol = "FFFFFF" if dark else NEUT_DARK
    tx_left = x+PAD+DN+0.12
    text(slide, tx_left, y0+0.06, (x+w-PAD)-tx_left, DN+0.20,
         [([(s, {'size': size, 'bold': True, 'color': tcol})],
           {'align': PP_ALIGN.LEFT, 'line_spacing': 0.98}) for s in title_lines],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)


def build():
    # Round 6 (Jose/Katy, 2026-06-25): the in-figure "Experimental Procedure"
    # title is REMOVED -- Katy cropped it out because the journal Figure 1
    # caption ("Overview of the Experimental Procedure") already supplies a
    # nearly identical header, so a 2nd one above the flow chart is redundant.
    # Canvas height tightened to the top of the cards now that the title row is
    # gone. Control chip also gains "may" ("...organizers may search on their
    # own") -- we do not KNOW organizers will search, so "may" is more accurate.
    y0 = 0.16; NH = 3.38
    W, H = 13.33, y0 + NH + y0
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _solid(bg, "FFFFFF"); _line(bg, None); _no_shadow(bg)

    ML = 0.40; UW = W - 2*ML
    g1 = 0.30; g2 = 0.58; g3 = 0.30
    widths = [2.76, 2.26, 3.62, 2.71]
    xs = []; cur = ML
    for i, wid in enumerate(widths):
        xs.append(cur); cur += wid + (g1, g2, g3, 0)[i]

    body_top = y0 + 0.74
    conn_cy = body_top + (NH - 0.74) / 2
    cap_top = body_top + 0.94

    def step(i, glyph):
        x = xs[i]; w = widths[i]
        rrect(slide, x, y0, w, NH, CARD, NEUT_LINE, 0.75, radius=0.05, shadow=True)
        badge(slide, x+w/2, body_top, glyph)
        return x, w

    def caption(x, w, runs, *, bottom):
        text(slide, x+0.20, cap_top, w-0.40, bottom-cap_top,
             [([runs], {'align': PP_ALIGN.CENTER, 'line_spacing': 1.05})],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 1 Email Outreach
    x, w = step(0, G_MAIL)
    header(slide, x, y0, w, 1, ["Email Outreach"])
    caption(x, w, ("1,586 seminar organizers and department chairs responsible for "
                   "1,881 STEM seminar series receive a link to a directory and "
                   "encouragement to diversify their speakers",
                   {'size': CAP_SZ, 'color': TX_DARK}), bottom=y0+NH-0.14)

    # 2 Directory Access
    x, w = step(1, G_KEY)
    header(slide, x, y0, w, 2, ["Directory Access"])
    caption(x, w, ("Each recipient enters a university email address to open the "
                   "directory built for their department",
                   {'size': CAP_SZ, 'color': TX_DARK}), bottom=y0+NH-0.14)

    # 3 Departments Randomly Assigned to Directory Type (red/blue chips)
    x = xs[2]; w = widths[2]
    rrect(slide, x, y0, w, NH, CARD, NEUT_LINE, 0.75, radius=0.05, shadow=True)
    header(slide, x, y0, w, 3, ["Departments Randomly Assigned", "to Directory Type"], size=12)

    chip_x = x + PAD; chip_w = w - 2*PAD
    chip_h = 1.12; chip_gap = 0.18
    t_top = conn_cy - chip_gap/2 - chip_h
    c_top = conn_cy + chip_gap/2

    def chip(ay, fill, accent, border, glyph, tag, label):
        rrect(slide, chip_x, ay, chip_w, chip_h, fill, line=border, line_w=0.9, radius=0.08)
        bd = 0.58; oval(slide, chip_x+0.20, ay+(chip_h-bd)/2, bd, bd, accent, line=None)
        text(slide, chip_x+0.20, ay+(chip_h-bd)/2, bd, bd,
             [[(glyph, {'font': ICON, 'size': 20, 'color': "FFFFFF"})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, chip_x+0.92, ay+0.10, chip_w-1.06, chip_h-0.20,
             [([(tag, {'size': 11, 'bold': True, 'color': accent})], {'space_after': 3}),
              ([(label, {'size': 10.5, 'color': TX_DARK})], {'line_spacing': 1.0})],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        return ay + chip_h/2

    t_cy = chip(t_top, TRT_FILL, TRT, TRT_BORDER, G_DIR, "TREATMENT",
                "Curated directory of qualified underrepresented racial minority (URM) faculty at peer departments")
    c_cy = chip(c_top, CTL_FILL, CTL, CTL_BORDER, G_LIST, "CONTROL",
                "Links to peer-department websites that organizers may search on their own")

    # 4 Outcome Measurement (softened dark endpoint card)
    x = xs[3]; w = widths[3]
    rrect(slide, x, y0, w, NH, CARD_DARK, line=None, radius=0.05, shadow=True)
    badge(slide, x+w/2, body_top, G_PEOPLE, fill="FFFFFF", gcolor=CARD_DARK, shadow=False)
    header(slide, x, y0, w, 4, ["Outcome", "Measurement"], dark=True,
           num_fill="FFFFFF", num_tcol=CARD_DARK)
    text(slide, x+0.20, cap_top, w-0.40, (y0+NH-0.52)-cap_top,
         [([("Share of presenting speakers who were URM, and separately Black and "
             "Hispanic, in the 1,686 continuing seminar series",
             {'size': CAP_SZ, 'color': "EDEFF2"})], {'align': PP_ALIGN.CENTER, 'line_spacing': 1.05})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    pill(slide, x+w/2, y0+NH-0.44, "2024-2025 academic year", fill="FFFFFF", tx=CARD_DARK)

    # ── connectors ───────────────────────────────────────────────────
    for a, b in ((0, 1), (2, 3)):
        rgt = xs[a]+widths[a]; lft = xs[b]
        chevron(slide, (rgt+lft)/2, conn_cy)

    # 2 -> 3 : RANDOMIZATION (in-spirit of Katy's shuffle). Style switchable.
    g_l = xs[1]+widths[1]; cy = conn_cy
    ex1 = chip_x - 0.03; gw = ex1 - g_l

    def _node_arms(nx):
        """short feed-in from step 2 + two arrows fanning to the two chips."""
        seg(slide, g_l+0.03, cy, nx, cy, FORK, 1.8)
        seg(slide, nx, cy, ex1-0.02, t_cy, FORK, 1.8, arrow=True)
        seg(slide, nx, cy, ex1-0.02, c_cy, FORK, 1.8, arrow=True)

    if RAND_STYLE == "shuffle":
        s = 0.15; sx0 = g_l + 0.04; stub = 0.06; gw2 = ex1 - sx0
        curve_arrow(slide, [(sx0, cy+s)] + _bez(
            (sx0+stub, cy+s), (sx0+0.20*gw2, cy+s), (sx0+0.72*gw2, t_cy), (ex1, t_cy)), FORK, 2.0)
        curve_arrow(slide, [(sx0, cy-s)] + _bez(
            (sx0+stub, cy-s), (sx0+0.20*gw2, cy-s), (sx0+0.72*gw2, c_cy), (ex1, c_cy)), FORK, 2.0)
    elif RAND_STYLE == "dice":
        nx = g_l + 0.30; _node_arms(nx)
        D = 0.42; x0 = nx-D/2; y0 = cy-D/2
        rrect(slide, x0, y0, D, D, "FFFFFF", line=NEUT_MED, line_w=1.4, radius=0.20, shadow=True)
        pr = 0.035
        for fx, fy in [(0.28, 0.28), (0.72, 0.28), (0.5, 0.5), (0.28, 0.72), (0.72, 0.72)]:
            oval(slide, x0+fx*D-pr, y0+fy*D-pr, 2*pr, 2*pr, NEUT_DARK, line=None)
    elif RAND_STYLE == "rnode":
        nx = g_l + 0.30; _node_arms(nx)
        d = 0.50; o = oval(slide, nx-d/2, cy-d/2, d, d, NEUT_MED, line="FFFFFF", line_w=1.5)
        soft_shadow(o, blur=0.045, dist=0.025, color="9AA2AC", alpha_pct=32)
        text(slide, nx-d/2, cy-d/2, d, d, [[("R", {'size': 17, 'bold': True, 'color': "FFFFFF"})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    out = OUT / "figure_1.pptx"; prs.save(str(out)); print("saved", out)


if __name__ == "__main__":
    build()
